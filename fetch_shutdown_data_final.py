"""
Fetch US Government Shutdown Financial Data from FRED using API
获取美国政府停摆期间的金融数据

Required: FRED API Key (free)
需要：FRED API密钥（免费）
Get it at: https://fred.stlouisfed.org/docs/api/api_key.html

Usage:
    python fetch_shutdown_data_final.py YOUR_API_KEY
"""

import sys
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import requests
import time

# Data series configuration
SERIES = {
    'WTREGEN': 'Treasury General Account (TGA)',
    'EFFR': 'Effective Federal Funds Rate (EFFR)',
    'SOFR': 'Secured Overnight Financing Rate (SOFR)',
    'WRESBAL': 'Bank Reserves'
}

# Shutdown periods with data range (1 year before to 1 year after)
SHUTDOWN_PERIODS = [
    {
        'name': '2013 Government Shutdown',
        'shutdown_start': '2013-10-01',
        'shutdown_end': '2013-10-17',
        'data_start': '2012-10-01',
        'data_end': '2014-10-17',
        'description': '16天停摆 (2013年10月1日至10月17日)'
    },
    {
        'name': '2018-2019 Government Shutdown',
        'shutdown_start': '2018-12-22',
        'shutdown_end': '2019-01-25',
        'data_start': '2017-12-22',
        'data_end': '2020-01-25',
        'description': '35天停摆 (2018年12月22日至2019年1月25日)'
    },
    {
        'name': '2025 Government Shutdown',
        'shutdown_start': '2025-10-01',
        'shutdown_end': None,
        'data_start': '2024-10-01',
        'data_end': datetime.now().strftime('%Y-%m-%d'),
        'description': '持续中 (2025年10月1日至今)'
    }
]


def fetch_fred_series_api(series_id, start_date, end_date, api_key):
    """Fetch data from FRED using official API"""
    print(f"  Fetching {series_id}...", end=' ')

    url = "https://api.stlouisfed.org/fred/series/observations"
    params = {
        'series_id': series_id,
        'api_key': api_key,
        'file_type': 'json',
        'observation_start': start_date,
        'observation_end': end_date
    }

    try:
        response = requests.get(url, params=params, timeout=30)
        response.raise_for_status()
        data = response.json()

        if 'observations' in data and len(data['observations']) > 0:
            records = []
            for obs in data['observations']:
                try:
                    value = float(obs['value']) if obs['value'] != '.' else None
                    records.append({
                        'date': obs['date'],
                        'value': value
                    })
                except (ValueError, KeyError):
                    continue

            if records:
                df = pd.DataFrame(records)
                df['date'] = pd.to_datetime(df['date'])
                df.set_index('date', inplace=True)
                df.columns = [series_id]
                print(f"✓ {len(df)} records")
                return df
            else:
                print(f"✗ No valid data")
                return None
        else:
            print(f"✗ No observations")
            return None

    except requests.exceptions.RequestException as e:
        print(f"✗ Request error: {e}")
        return None
    except Exception as e:
        print(f"✗ Error: {e}")
        return None


def fetch_all_data_for_period(period, api_key):
    """Fetch all series for a given period"""
    print(f"\n{'=' * 70}")
    print(f"{period['name']}")
    print(f"{period['description']}")
    print(f"Data range: {period['data_start']} to {period['data_end']}")
    print(f"{'=' * 70}")

    all_data = None

    for series_id, series_name in SERIES.items():
        df = fetch_fred_series_api(series_id, period['data_start'], period['data_end'], api_key)

        if df is not None and not df.empty:
            if all_data is None:
                all_data = df
            else:
                all_data = all_data.join(df, how='outer')

        time.sleep(0.2)  # Rate limiting

    if all_data is not None:
        all_data.sort_index(inplace=True)
        print(f"\n✓ Total records: {len(all_data)}")

    return all_data


def add_shutdown_highlights(ws, period, start_row):
    """Add visual highlights for shutdown periods"""
    shutdown_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")

    if period['shutdown_end']:
        shutdown_start = pd.to_datetime(period['shutdown_start'])
        shutdown_end = pd.to_datetime(period['shutdown_end'])

        # Highlight rows within shutdown period
        for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row):
            try:
                date_cell = row[0]
                if isinstance(date_cell.value, datetime):
                    if shutdown_start <= date_cell.value <= shutdown_end:
                        for cell in row:
                            cell.fill = shutdown_fill
            except:
                continue


def create_excel_file(data_dict, filename='us_shutdown_financial_data.xlsx'):
    """Create formatted Excel file with multiple sheets"""
    print(f"\n{'=' * 70}")
    print(f"Creating Excel file: {filename}")
    print(f"{'=' * 70}")

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        # Summary sheet
        summary_data = []
        for i, period in enumerate(SHUTDOWN_PERIODS):
            period_name = period['name']
            if period_name in data_dict:
                df = data_dict[period_name]
                summary_data.append({
                    'No.': i + 1,
                    'Period': period['description'],
                    'Data Start': df.index.min().strftime('%Y-%m-%d'),
                    'Data End': df.index.max().strftime('%Y-%m-%d'),
                    'Total Records': len(df),
                    'Indicators': ', '.join([SERIES[col] for col in df.columns])
                })

        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name='Summary', index=False)

        # Data sheets for each period
        for i, period in enumerate(SHUTDOWN_PERIODS):
            period_name = period['name']
            if period_name in data_dict:
                df = data_dict[period_name]
                # Rename columns to full names
                df_export = df.copy()
                df_export.columns = [SERIES[col] for col in df.columns]

                sheet_name = f"{i+1}_{period['name'].split()[0]}"
                df_export.to_excel(writer, sheet_name=sheet_name)

    # Format the workbook
    wb = openpyxl.load_workbook(filename)

    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Style header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
                except:
                    pass
            adjusted_width = min(max_length + 3, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Freeze top row
        ws.freeze_panes = 'A2'

        # Add borders to all cells with data
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                if cell.value:
                    cell.border = border

        # Highlight shutdown periods in data sheets
        if sheet_name != 'Summary':
            for period in SHUTDOWN_PERIODS:
                if period['name'].split()[0] in sheet_name:
                    add_shutdown_highlights(ws, period, 2)

    wb.save(filename)
    print(f"✓ Excel file created successfully!")
    print(f"  File: {filename}")
    print(f"  Sheets: {len(wb.sheetnames)}")


def main():
    """Main function"""
    print("\n" + "=" * 70)
    print("US Government Shutdown Financial Data Fetcher")
    print("美国政府停摆金融数据获取工具")
    print("=" * 70)

    # Check for API key
    if len(sys.argv) < 2:
        print("\n✗ Error: FRED API key required")
        print("\nUsage:")
        print(f"  python {sys.argv[0]} YOUR_API_KEY")
        print("\nGet a free API key at:")
        print("  https://fred.stlouisfed.org/docs/api/api_key.html")
        print("\nExample:")
        print(f"  python {sys.argv[0]} abcdef1234567890abcdef1234567890")
        sys.exit(1)

    api_key = sys.argv[1]
    print(f"\n✓ Using API key: {api_key[:10]}...")

    # Fetch data for all periods
    all_period_data = {}

    for period in SHUTDOWN_PERIODS:
        df = fetch_all_data_for_period(period, api_key)
        if df is not None and not df.empty:
            all_period_data[period['name']] = df
        else:
            print(f"⚠ Warning: No data fetched for {period['name']}")

    # Create Excel file
    if all_period_data:
        create_excel_file(all_period_data)

        print("\n" + "=" * 70)
        print("✓ SUCCESS! Data fetched and saved")
        print("=" * 70)

        print("\nData Summary:")
        for period_name, df in all_period_data.items():
            print(f"\n{period_name}:")
            print(f"  Records: {len(df)}")
            print(f"  Date range: {df.index.min().date()} to {df.index.max().date()}")
            print(f"  Indicators: {len(df.columns)}")
            for col in df.columns:
                non_null = df[col].count()
                print(f"    - {SERIES[col]}: {non_null} values")

        print(f"\n✓ Excel file: us_shutdown_financial_data.xlsx")
    else:
        print("\n✗ Failed to fetch any data")
        print("Please check your API key and internet connection")


if __name__ == "__main__":
    main()
